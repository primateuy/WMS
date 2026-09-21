# -*- coding: utf-8 -*-
"""Lectura del Excel de operaciones a una lista de filas normalizadas.

Los ``.xlsx`` se leen con ``openpyxl`` y los ``.xls`` con ``xlrd``. No se
puede depender de ``xlrd`` para ambos: desde su versión 2.0 dejó de soportar
``.xlsx`` ("Excel xlsx file; not supported"), y según el servidor puede haber
instalada la 1.2 o la 2.x. El formato se detecta por la firma binaria del
archivo, no por la extensión. La conversión de celdas replica la de
``base_import``: todo se entrega como texto ya recortado, las fechas en
formato ISO y los enteros sin ".0".
"""
import datetime
import io
import logging

from odoo import _
from odoo.exceptions import UserError
from odoo.tools import DEFAULT_SERVER_DATE_FORMAT, DEFAULT_SERVER_DATETIME_FORMAT

try:
    import openpyxl
except ImportError:
    openpyxl = None
try:
    import xlrd
except ImportError:
    xlrd = None

_logger = logging.getLogger(__name__)

# Columnas de la plantilla, en el orden en que se generan.
COLUMNS = (
    'operation_ref',
    'picking_type',
    'warehouse',
    'partner',
    'location_origin',
    'location_dest',
    'scheduled_date',
    'origin',
    'product',
    'lot',
    'quantity',
    'uom',
)
REQUIRED_COLUMNS = ('operation_ref', 'picking_type', 'product', 'quantity')

# Firmas binarias: .xlsx es un ZIP y .xls un contenedor OLE2.
XLSX_SIGNATURE = b'PK\x03\x04'
XLS_SIGNATURE = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'

# Columnas que definen el cabezal de la operación: deben coincidir en todas las
# filas con la misma operation_ref.
HEADER_COLUMNS = (
    'picking_type',
    'warehouse',
    'partner',
    'location_origin',
    'location_dest',
    'scheduled_date',
    'origin',
)


class ExcelParser:
    """Convierte el archivo Excel en filas ``dict`` listas para validar."""

    def __init__(self, env):
        self.env = env

    def parse(self, file_bytes):
        """Lee la primera hoja del archivo y devuelve sus filas.

        Args:
            file_bytes (bytes): contenido binario del .xlsx/.xls.

        Returns:
            list[dict]: una entrada por fila con datos. Cada dict tiene la
                clave ``row`` (número de fila en Excel, base 1) y una clave por
                columna de ``COLUMNS`` con el valor como texto ('' si vacío).

        Raises:
            UserError: si falta la librería, el archivo no se puede abrir, no
                tiene encabezado o faltan columnas obligatorias.
        """
        raw_rows = self._read_rows(file_bytes or b'')
        if not raw_rows:
            raise UserError(_("El archivo está vacío: no tiene fila de encabezados."))

        column_index = self._read_header(raw_rows[0])
        rows = []
        for position, raw_row in enumerate(raw_rows[1:], start=2):
            values = {
                column: raw_row[index] if index is not None and index < len(raw_row) else ''
                for column, index in column_index.items()
            }
            if not any(values.values()):
                continue
            values['row'] = position
            rows.append(values)
        if not rows:
            raise UserError(_("El archivo no tiene filas de datos debajo del encabezado."))
        return rows

    def _read_rows(self, file_bytes):
        """Devuelve la primera hoja como lista de filas de texto.

        Elige la librería según la firma binaria del archivo: ``openpyxl``
        para .xlsx y ``xlrd`` para .xls.

        Returns:
            list[list[str]]: filas (incluido el encabezado) con cada celda ya
                convertida a texto.

        Raises:
            UserError: si el formato no se reconoce, falta la librería
                correspondiente o el archivo no se puede abrir.
        """
        if file_bytes.startswith(XLSX_SIGNATURE):
            reader = self._read_rows_xlsx
        elif file_bytes.startswith(XLS_SIGNATURE):
            reader = self._read_rows_xls
        else:
            raise UserError(_(
                "No se pudo leer el archivo. Verificá que sea un Excel válido (.xlsx o .xls)."
            ))
        try:
            return reader(file_bytes)
        except UserError:
            raise
        except Exception as error:  # cada librería lanza sus propios tipos según el formato
            _logger.info("No se pudo abrir el Excel de importación: %s", error)
            raise UserError(_(
                "No se pudo leer el archivo. Verificá que sea un Excel válido (.xlsx o .xls). "
                "Detalle: %s", error
            )) from error

    def _read_rows_xlsx(self, file_bytes):
        """Lee un .xlsx con openpyxl (valores calculados, sin fórmulas)."""
        if openpyxl is None:
            raise UserError(_(
                "Falta la librería Python 'openpyxl' en el servidor, necesaria para leer "
                "archivos .xlsx. Instalala con: pip install openpyxl"
            ))
        book = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        try:
            sheet = book.worksheets[0]
            return [
                [self._value_to_text(value) for value in row]
                for row in sheet.iter_rows(values_only=True)
            ]
        finally:
            book.close()

    def _read_rows_xls(self, file_bytes):
        """Lee un .xls con xlrd (cualquier versión soporta este formato)."""
        if xlrd is None:
            raise UserError(_(
                "Falta la librería Python 'xlrd' en el servidor, necesaria para leer "
                "archivos .xls. Instalala con: pip install xlrd"
            ))
        book = xlrd.open_workbook(file_contents=file_bytes)
        sheet = book.sheet_by_index(0)
        return [
            [self._cell_to_text(book, sheet.cell(row_number, index)) for index in range(sheet.ncols)]
            for row_number in range(sheet.nrows)
        ]

    def _read_header(self, header_row):
        """Mapea cada columna conocida a su índice en la fila 1.

        Los encabezados se comparan en minúsculas y sin espacios en los bordes.
        Columnas opcionales ausentes quedan mapeadas a ``None``.

        Returns:
            dict: {nombre_columna: índice | None}.

        Raises:
            UserError: si falta alguna columna de ``REQUIRED_COLUMNS``.
        """
        headers = [header.lower() for header in header_row]
        column_index = {
            column: headers.index(column) if column in headers else None
            for column in COLUMNS
        }
        missing = [column for column in REQUIRED_COLUMNS if column_index[column] is None]
        if missing:
            raise UserError(_(
                "Faltan columnas obligatorias en el encabezado: %s. "
                "Descargá la plantilla desde el asistente para ver el formato esperado.",
                ", ".join(missing),
            ))
        return column_index

    @staticmethod
    def _value_to_text(value):
        """Devuelve un valor tipado de openpyxl como texto recortado.

        Misma salida que ``_cell_to_text``: enteros sin ".0", fechas en
        formato ISO (con hora si la tienen) y booleanos como 'True'/'False'.
        """
        if value is None:
            return ''
        if isinstance(value, bool):
            return 'True' if value else 'False'
        if isinstance(value, int):
            return str(value)
        if isinstance(value, float):
            return str(int(value)) if value % 1 == 0.0 else repr(value)
        if isinstance(value, datetime.datetime):
            has_time = (value.hour, value.minute, value.second) != (0, 0, 0)
            return value.strftime(
                DEFAULT_SERVER_DATETIME_FORMAT if has_time else DEFAULT_SERVER_DATE_FORMAT
            )
        if isinstance(value, datetime.date):
            return value.strftime(DEFAULT_SERVER_DATE_FORMAT)
        return str(value).strip()

    @staticmethod
    def _cell_to_text(book, cell):
        """Devuelve el valor de una celda de xlrd como texto recortado.

        Números enteros sin decimales, fechas en formato ISO (con hora si la
        celda la tiene) y booleanos como 'True'/'False'.
        """
        if cell.ctype == xlrd.XL_CELL_NUMBER:
            value = cell.value
            return str(int(value)) if value % 1 == 0.0 else repr(value)
        if cell.ctype == xlrd.XL_CELL_DATE:
            date_value = datetime.datetime(*xlrd.xldate.xldate_as_tuple(cell.value, book.datemode))
            has_time = cell.value % 1 != 0.0
            return date_value.strftime(
                DEFAULT_SERVER_DATETIME_FORMAT if has_time else DEFAULT_SERVER_DATE_FORMAT
            )
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return 'True' if cell.value else 'False'
        if cell.ctype == xlrd.XL_CELL_ERROR:
            return ''
        return str(cell.value or '').strip()
